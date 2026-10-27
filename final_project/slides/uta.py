"""Helpers for building slides on the official UTA PowerPoint template (python-pptx).

The template is downloaded from UTA's brand site on first use. Brand rules followed here:
UTA Blue #0064B1 is the primary colour, UTA Orange #F58025 is used for accents only (never for
text), titles use the template's own title/subtitle placeholders.
"""
import urllib.request
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
TEMPLATE_URL = ("https://resources.uta.edu/mme/identity/_files/"
                "UTA%20PPT%20Instructions%20and%20Template%2011-08-22.pptx")
TEMPLATE = HERE / ".cache" / "uta_template.pptx"
ICONS = HERE / "assets" / "icons"

BLUE, BLUE_TXT, NAVY, ORANGE = "0064B1", "00599B", "003865", "F58025"
LIGHT, LIGHT2, SKY = "EDF3F9", "BDD6E6", "7BAFD4"
GRAY, GRAY_L, HAIR, TEXT, WHITE = "5B6770", "A2AAAD", "D0D3D4", "1F2933", "FFFFFF"
FONT, MONO = "Arial", "Courier New"
TOP, BOT = 1.24, 4.55          # usable content band on the template's content layout (inches)


def rgb(h):
    return RGBColor.from_string(h)


def template():
    if not TEMPLATE.exists():
        TEMPLATE.parent.mkdir(parents=True, exist_ok=True)
        urllib.request.urlretrieve(TEMPLATE_URL, TEMPLATE)
    prs = Presentation(TEMPLATE)
    lst = prs.slides._sldIdLst
    for sid in list(lst):            # drop the instruction/demo slides, keep masters and layouts
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)
    return prs


def mark_decorative(shape):
    c_nv_pr = shape._element.find(".//" + qn("p:cNvPr"))
    ext_lst = c_nv_pr.find(qn("a:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(c_nv_pr, qn("a:extLst"))
    ext = etree.SubElement(ext_lst, qn("a:ext"))
    ext.set("uri", "{C183D7F6-B498-43B3-948B-1728B52AA6E4}")
    adec = "http://schemas.microsoft.com/office/drawing/2017/decorative"
    etree.SubElement(ext, "{%s}decorative" % adec, nsmap={"adec": adec}).set("val", "1")


def text(slide, x, y, w, h, paras, size=10, color=TEXT, bold=False, align="l", anchor="t",
         margin=(0, 0, 0, 0), space_after=0, spacing=None, name=None):
    """paras: list of str or list of (text, opts) runs; opts: bold, color, size, font, italic."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margin]
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if spacing:
            p.line_spacing = spacing
        if i < len(paras) - 1 and space_after:
            p.space_after = Pt(space_after)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", FONT)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = rgb(o.get("color", color))
    if name:
        tb.name = name
    return tb


def box(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=None, decorative=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = radius / min(w, h)
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if decorative:
        mark_decorative(shp)
    return shp


def oval(slide, cx, cy, d, fill, line=None, line_w=1.5):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def badge(slide, cx, cy, d, label, fill=NAVY, ring=ORANGE, size=9, color=WHITE):
    shp = oval(slide, cx, cy, d, fill, ring, 1.5)
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size, r.font.bold, r.font.name = Pt(size), True, FONT
    r.font.color.rgb = rgb(color)
    return shp


def icon(slide, cx, cy, d, name, fill=BLUE, scale=0.56):
    circ = oval(slide, cx, cy, d, fill)
    mark_decorative(circ)
    s = d * scale
    pic = slide.shapes.add_picture(str(ICONS / f"{name}.png"), Inches(cx - s / 2), Inches(cy - s / 2), Inches(s), Inches(s))
    pic._element.find(".//" + qn("p:cNvPr")).set("descr", "")
    mark_decorative(pic)
    return circ


def line(slide, x1, y1, x2, y2, color, width=1.5):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(width)
    mark_decorative(ln)
    return ln


def card_header(slide, x, y, w, icon_name, title, d=0.32, size=11, color=NAVY, icon_fill=BLUE):
    icon(slide, x + 0.14 + d / 2, y + 0.09 + d / 2, d, icon_name, fill=icon_fill)
    text(slide, x + 0.24 + d, y + 0.09, w - d - 0.34, d, [title], size=size, bold=True, color=color, anchor="m")


def picture(slide, path, x, y, w=None, h=None, alt=""):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    pic._element.find(".//" + qn("p:cNvPr")).set("descr", alt)
    return pic


def set_ph(ph, value, size=None, bold=None, color=None):
    ph.text_frame.text = value
    for r in ph.text_frame.paragraphs[0].runs:
        if size:
            r.font.size = Pt(size)
        if bold is not None:
            r.font.bold = bold
        if color:
            r.font.color.rgb = rgb(color)


def place(ph, x, y, w, h):
    ph.left, ph.top, ph.width, ph.height = Inches(x), Inches(y), Inches(w), Inches(h)


def content_slide(prs, title, subtitle, notes=None):
    layout = prs.slide_layouts[5]
    assert layout.name == "Content Slide"
    s = prs.slides.add_slide(layout)
    place(s.shapes.title, 0.5, 0.14, 9.0, 0.66)
    set_ph(s.shapes.title, title, size=28)
    sub = s.placeholders[10]
    place(sub, 0.5, 0.76, 9.0, 0.34)
    set_ph(sub, subtitle, size=14)
    body = s.placeholders[1]
    body._element.getparent().remove(body._element)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def chart_clean(chart, size=8):
    cs = chart._chartSpace
    ns = ('xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
          'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"')
    old = cs.find(qn("c:spPr"))
    if old is not None:
        cs.remove(old)
    cs.find(qn("c:chart")).addnext(etree.fromstring(f"<c:spPr {ns}><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>"))
    plot_area = cs.find(".//" + qn("c:plotArea"))
    sp2 = etree.fromstring(f"<c:spPr {ns}><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>")
    ext = plot_area.find(qn("c:extLst"))
    (ext.addprevious(sp2) if ext is not None else plot_area.append(sp2))
    chart.font.name = FONT
    chart.font.size = Pt(size)
    chart.font.color.rgb = rgb(GRAY)


def point_label(pt, label, color=NAVY, size=8, bold=True, pos=XL_LABEL_POSITION.OUTSIDE_END):
    dl = pt.data_label
    tf = dl.text_frame
    tf.text = label
    for r in tf.paragraphs[0].runs:
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = rgb(color)
    dl.position = pos


def set_alt(graphic_frame, alt):
    graphic_frame._element.find(".//" + qn("p:cNvPr")).set("descr", alt)


__all__ = [n for n in dir() if not n.startswith("_")] + ["CategoryChartData", "XL_LABEL_POSITION", "XL_TICK_MARK",
                                                         "Inches", "Pt"]
